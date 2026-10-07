"""Panel application for NeuroPipe-SABV.

Run locally:
    panel serve panel_app.py --show --autoreload
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import numpy as np
import pandas as pd
import panel as pn
import matplotlib.pyplot as plt
from scipy import stats as scipy_stats
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from src.cleaning import basic_clean
from src.features import add_log_feature

pn.extension("tabulator", sizing_mode="stretch_width")


ACCENT = "#22c55e"
BG = "#090b11"
PANEL = "#111524"
TEXT = "#e6ebf5"
MUTED = "#9ca7bf"

CSS = """
:root {
  --bg: #090b11;
  --panel: #111524;
  --soft: #191f33;
  --text: #e6ebf5;
  --muted: #9ca7bf;
  --accent: #22c55e;
  --line: #2b344c;
}
body, .bk, .pn-template {
  background: radial-gradient(circle at 18% -5%, #1f2b4f 0%, transparent 42%),
    radial-gradient(circle at 90% 0%, #19352a 0%, transparent 34%), var(--bg);
  color: var(--text);
}
.hero {
  background: linear-gradient(130deg, #0f1426 0%, #15243a 100%);
  border: 1px solid var(--line);
  padding: 24px;
  margin-bottom: 14px;
  box-shadow: 0 24px 48px rgba(0, 0, 0, 0.36);
}
.hero h1 {
  margin: 0 0 8px;
  font-size: 42px;
  line-height: 1.05;
  color: #f2f5ff;
}
.hero p, .muted {
  color: var(--muted);
}
.metric {
  background: rgba(255, 255, 255, 0.04);
  border: 1px solid var(--line);
  padding: 12px;
}
.metric b {
  display: block;
  color: #ffffff;
  font-size: 22px;
}
.metric span {
  color: var(--muted);
  font-size: 13px;
}
.section {
  border-top: 1px solid var(--line);
  padding-top: 14px;
  margin-top: 12px;
}
"""

pn.config.raw_css.append(CSS)


def parse_animal_number(value) -> float:
    """Extract animal ID numbers from numeric values or labels."""
    if pd.isna(value):
        return np.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    text = str(value).strip()
    if not text:
        return np.nan
    lower_text = text.lower()
    prefixed_match = re.search(
        r"(?:rat|subject|animal)\s*[-_#:]?\s*(\d+(?:\.\d+)?)",
        lower_text,
    )
    if prefixed_match:
        return float(prefixed_match.group(1))
    numeric_only_match = re.fullmatch(r"\d+(?:\.\d+)?", lower_text)
    if numeric_only_match:
        return float(numeric_only_match.group(0))
    match = re.search(r"\d+(?:\.\d+)?", text)
    if match:
        return float(match.group(0))
    return np.nan


def parse_animal_number_series(series: pd.Series) -> pd.Series:
    return series.apply(parse_animal_number)


def parse_id_list(raw_text: str):
    values = set()
    invalid_tokens = []
    for token in raw_text.split(","):
        token = token.strip()
        if not token:
            continue
        if "-" in token:
            parts = token.split("-", 1)
            try:
                start = float(parts[0].strip())
                end = float(parts[1].strip())
            except ValueError:
                invalid_tokens.append(token)
                continue
            if float(start).is_integer() and float(end).is_integer():
                lo, hi = sorted((int(start), int(end)))
                values.update(float(v) for v in range(lo, hi + 1))
            else:
                values.update((float(start), float(end)))
            continue
        try:
            values.add(float(token))
        except ValueError:
            invalid_tokens.append(token)
    return values, invalid_tokens


def demo_dataframe() -> pd.DataFrame:
    np.random.seed(42)
    n_subjects = 30
    ids = np.arange(1, n_subjects + 1)
    distance = np.where(
        ids <= 15,
        np.random.normal(3000, 500, n_subjects),
        np.random.normal(2500, 400, n_subjects),
    )
    velocity = distance / 300
    center_time = np.where(
        ids <= 15,
        np.random.normal(15, 5, n_subjects),
        np.random.normal(25, 6, n_subjects),
    )
    distance[2] = np.nan
    distance[28] = 9500.0
    velocity[10] = np.nan
    return pd.DataFrame(
        {
            "Subject_ID": [f"rat_{i}" for i in ids],
            "Distance_Moved_cm": distance,
            "Velocity_cm_s": velocity,
            "Center_Time_s": center_time,
            "Test_Date": ["2026-08-01"] * n_subjects,
        }
    )


def read_uploaded_file(value: bytes, filename: str) -> pd.DataFrame:
    suffix = Path(filename or "").suffix.lower()
    buffer = io.BytesIO(value)
    if suffix == ".csv":
        for encoding in ("utf-8", "latin1", "cp1252"):
            try:
                buffer.seek(0)
                return pd.read_csv(buffer, encoding=encoding, low_memory=False)
            except Exception:
                continue
        raise ValueError("Could not decode CSV. Try saving as an Excel file.")
    if suffix in {".xlsx", ".xls"}:
        return pd.read_excel(buffer, sheet_name=0)
    raise ValueError("Upload a CSV, XLSX, or XLS file.")


def ttest_for_groups(df: pd.DataFrame, value_col: str, group_col: str = "Sex"):
    male_data = df[df[group_col] == "Male"][value_col].dropna()
    female_data = df[df[group_col] == "Female"][value_col].dropna()
    if len(male_data) > 1 and len(female_data) > 1:
        t_stat, p_value = scipy_stats.ttest_ind(male_data, female_data, equal_var=False)
        return t_stat, p_value, len(male_data), len(female_data)
    return np.nan, np.nan, len(male_data), len(female_data)


def effect_size_cohens_d(group1: pd.Series, group2: pd.Series) -> float:
    n1, n2 = len(group1), len(group2)
    if n1 < 2 or n2 < 2:
        return np.nan
    pooled_std = np.sqrt(
        ((n1 - 1) * group1.var(ddof=1) + (n2 - 1) * group2.var(ddof=1)) / (n1 + n2 - 2)
    )
    if pooled_std == 0:
        return np.nan
    return (group1.mean() - group2.mean()) / pooled_std


def sem(series: pd.Series) -> float:
    values = series.dropna()
    if len(values) < 2:
        return np.nan
    return values.std(ddof=1) / np.sqrt(len(values))


def safe_percent_diff(male_mean: float, female_mean: float) -> float:
    if pd.isna(female_mean) or female_mean == 0:
        return np.nan
    return (male_mean - female_mean) / female_mean * 100


def infer_animal_column(df: pd.DataFrame):
    normalized = {col.lower().strip(): col for col in df.columns}
    preferred = ["animal #", "animal", "animal number", "animal_number", "animal id", "animal_id"]
    for key in preferred:
        if key in normalized:
            return normalized[key]
    for col in df.columns:
        if "animal" in col.lower() or "subject" in col.lower() or "rat" in col.lower():
            return col
    return None


def autosize_columns(ws, min_width: int = 10, max_width: int = 42) -> None:
    for col_cells in ws.columns:
        max_len = 0
        col_idx = col_cells[0].column
        for cell in col_cells:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, min_width), max_width)


def build_excel_export(
    df_processed: pd.DataFrame,
    animal_col: str | None = None,
    threshold: float = 16,
    rebuild_sex_labels: bool = True,
) -> bytes:
    excel_buffer = io.BytesIO()
    animal_col = animal_col if animal_col in df_processed.columns else infer_animal_column(df_processed)
    export_df = df_processed.copy()
    if animal_col and rebuild_sex_labels:
        animal_num = parse_animal_number_series(export_df[animal_col])
        export_df["Sex"] = np.where(
            animal_num <= threshold,
            "Male",
            np.where(animal_num > threshold, "Female", "Unclassified"),
        )
    elif "Sex" not in export_df.columns:
        export_df["Sex"] = "Unclassified"

    with pd.ExcelWriter(excel_buffer, engine="openpyxl") as writer:
        sheet_name = "Summary_By_Animal"
        ordered_df = export_df.copy()
        low_count = 0
        has_gap = False
        if animal_col:
            ordered_df["_animal_num_sort"] = parse_animal_number_series(ordered_df[animal_col])
            low = ordered_df[ordered_df["_animal_num_sort"].notna() & (ordered_df["_animal_num_sort"] <= threshold)]
            high = ordered_df[ordered_df["_animal_num_sort"].notna() & (ordered_df["_animal_num_sort"] > threshold)]
            unknown = ordered_df[ordered_df["_animal_num_sort"].isna()]
            low_count = len(low)
            has_gap = len(low) > 0 and len(high) > 0
            pieces = [low]
            if has_gap:
                pieces.append(pd.DataFrame([{col: np.nan for col in ordered_df.columns}]))
            pieces.extend([high, unknown])
            ordered_df = pd.concat(pieces, ignore_index=True).drop(columns=["_animal_num_sort"])

        ordered_df.to_excel(writer, index=False, sheet_name=sheet_name, startrow=2)
        ws = writer.sheets[sheet_name]
        header_font = Font(bold=True)
        section_font = Font(bold=True, size=12)
        section_fill = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
        table_header_fill = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
        highlight_fill = PatternFill(start_color="FFF59D", end_color="FFF59D", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        ws.cell(row=1, column=1, value="Summary Ordered by Animal Number").font = Font(bold=True, size=13)
        ws.cell(row=3, column=1, value="Animal-Ordered Processed Data").font = section_font
        ws.cell(row=3, column=1).fill = section_fill
        for c in range(1, len(ordered_df.columns) + 1):
            ws.cell(row=3, column=c).font = header_font
            ws.cell(row=3, column=c).fill = table_header_fill
            ws.cell(row=3, column=c).border = thin_border
        if animal_col and has_gap:
            spacer_row = 4 + low_count
            ws.cell(row=spacer_row, column=1, value=f"--- Spacer: animals > {threshold} begin below ---").font = header_font
            ws.cell(row=spacer_row, column=1).fill = section_fill

        summary_title_row = 4 + len(ordered_df) + 2
        summary_header_row = summary_title_row + 1
        summary_data_start = summary_header_row + 1
        summary_headers = [
            "Variable",
            "Male N",
            "Female N",
            "Male Mean",
            "Female Mean",
            "Male SEM",
            "Female SEM",
            "t-stat",
            "p-value",
            "p<0.005",
        ]
        ws.cell(row=summary_title_row, column=1, value="Sex Summary Table (Welch t-test)").font = Font(bold=True, size=12)
        for i, col_title in enumerate(summary_headers, start=1):
            cell = ws.cell(row=summary_header_row, column=i, value=col_title)
            cell.font = header_font
            cell.fill = table_header_fill
            cell.border = thin_border

        numeric_cols = [col for col in export_df.select_dtypes(include=["number"]).columns if col != animal_col]
        for i, col_name in enumerate(numeric_cols):
            row = summary_data_start + i
            male_vals = export_df[export_df["Sex"] == "Male"][col_name].dropna()
            female_vals = export_df[export_df["Sex"] == "Female"][col_name].dropna()
            t_stat, p_value, n_male, n_female = ttest_for_groups(export_df, col_name)
            values = [
                col_name,
                n_male,
                n_female,
                male_vals.mean(),
                female_vals.mean(),
                sem(male_vals),
                sem(female_vals),
                t_stat,
                p_value,
                "YES" if not pd.isna(p_value) and p_value < 0.005 else "NO",
            ]
            for col_num, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col_num, value=value)
                cell.border = thin_border
                if values[-1] == "YES":
                    cell.fill = highlight_fill
        ws.freeze_panes = "A4"
        autosize_columns(ws)
    return excel_buffer.getvalue()


class NeuroPipePanel:
    def __init__(self):
        self.df_raw: pd.DataFrame | None = None
        self.df_processed: pd.DataFrame | None = None
        self.uploaded_filename = "processed_data"
        self.sex_col_used = None
        self.sex_threshold_used = 16.0
        self.sex_rebuild_from_threshold = True

        self.file_input = pn.widgets.FileInput(accept=".csv,.xlsx,.xls", name="Upload CSV or Excel")
        self.demo_button = pn.widgets.Button(name="Load Demo Data", button_type="primary")
        self.rows_slider = pn.widgets.IntSlider(name="Rows to display", start=5, end=50, value=10)

        self.sex_col = pn.widgets.Select(name="Column for sex classification", options=[])
        self.mode = pn.widgets.RadioButtonGroup(
            name="Classification method",
            options=["Threshold split", "Manual number lists", "Female list (others Male)"],
            value="Threshold split",
            button_type="success",
        )
        self.threshold = pn.widgets.IntInput(name="Threshold", value=16, step=1)
        self.male_ids = pn.widgets.TextInput(name="Male IDs", value="1-16")
        self.female_ids = pn.widgets.TextInput(name="Female IDs", value="17-32")
        self.process_button = pn.widgets.Button(name="Run Processing Pipeline", button_type="primary")
        self.compare_var = pn.widgets.Select(name="Variable to compare", options=[])
        self.multi_vars = pn.widgets.MultiChoice(name="Variables for summary", options=[], value=[])

        self.status = pn.pane.Alert("Upload data or load the demo dataset to begin.", alert_type="info")
        self.upload_view = pn.Column()
        self.process_view = pn.Column()
        self.analysis_view = pn.Column()
        self.download_view = pn.Column()

        self.file_input.param.watch(self._load_uploaded, "value")
        self.demo_button.on_click(self._load_demo)
        self.process_button.on_click(self._process_data)
        self.compare_var.param.watch(lambda _event: self._render_analysis(), "value")
        self.multi_vars.param.watch(lambda _event: self._render_analysis(), "value")

        self._render_all()

    def _load_uploaded(self, event):
        if not event.new:
            return
        try:
            self.df_raw = read_uploaded_file(event.new, self.file_input.filename)
            self.uploaded_filename = self.file_input.filename or "uploaded_data"
            self.df_processed = None
            self._after_raw_load(f"Loaded {self.uploaded_filename}: {len(self.df_raw):,} rows x {len(self.df_raw.columns)} columns.")
        except Exception as exc:
            self.status.object = f"Error loading file: {exc}"
            self.status.alert_type = "danger"

    def _load_demo(self, _event):
        self.df_raw = demo_dataframe()
        self.uploaded_filename = "demo_open_field_data.csv"
        self.df_processed = None
        self._after_raw_load("Loaded demo open-field dataset.")

    def _after_raw_load(self, message: str):
        options = list(self.df_raw.columns)
        self.sex_col.options = options
        self.sex_col.value = options[0] if options else None
        self.status.object = message
        self.status.alert_type = "success"
        self._render_all()

    def _process_data(self, _event):
        if self.df_raw is None or not self.sex_col.value:
            self.status.object = "Load data and select an ID column before processing."
            self.status.alert_type = "warning"
            return
        try:
            df_processed, clean_report = basic_clean(self.df_raw)
            actual_sex_col = self.sex_col.value.strip() if isinstance(self.sex_col.value, str) else self.sex_col.value
            renamed_map = clean_report.get("header_standardization", {}).get("renamed", {})
            if isinstance(renamed_map, dict):
                actual_sex_col = renamed_map.get(actual_sex_col, actual_sex_col)
            sex_numeric = parse_animal_number_series(df_processed[actual_sex_col])

            if self.mode.value == "Threshold split":
                threshold = float(self.threshold.value)
                df_processed["Sex"] = np.where(
                    sex_numeric <= threshold,
                    "Male",
                    np.where(sex_numeric > threshold, "Female", "Unclassified"),
                )
                rebuild = True
            elif self.mode.value == "Manual number lists":
                male_ids, male_invalid = parse_id_list(self.male_ids.value)
                female_ids, female_invalid = parse_id_list(self.female_ids.value)
                invalid = male_invalid + female_invalid
                if invalid:
                    raise ValueError(f"Invalid ID tokens: {', '.join(invalid)}")
                df_processed["Sex"] = np.select(
                    [sex_numeric.isin(male_ids), sex_numeric.isin(female_ids)],
                    ["Male", "Female"],
                    default="Unclassified",
                )
                rebuild = False
            else:
                female_ids, invalid = parse_id_list(self.female_ids.value)
                if invalid:
                    raise ValueError(f"Invalid ID tokens: {', '.join(invalid)}")
                df_processed["Sex"] = np.select(
                    [sex_numeric.isin(female_ids), sex_numeric.notna()],
                    ["Female", "Male"],
                    default="Unclassified",
                )
                rebuild = False

            numeric_cols = df_processed.select_dtypes(include=["number"]).columns
            if len(numeric_cols) > 0:
                df_processed = add_log_feature(df_processed, col=numeric_cols[0])

            self.df_processed = df_processed
            self.sex_col_used = actual_sex_col
            self.sex_threshold_used = float(self.threshold.value)
            self.sex_rebuild_from_threshold = rebuild
            numeric_options = df_processed.select_dtypes(include=["number"]).columns.tolist()
            self.compare_var.options = numeric_options
            self.compare_var.value = numeric_options[0] if numeric_options else None
            self.multi_vars.options = numeric_options
            self.multi_vars.value = numeric_options[: min(3, len(numeric_options))]
            unclassified = int((df_processed["Sex"] == "Unclassified").sum())
            suffix = f" {unclassified} rows are unclassified." if unclassified else ""
            self.status.object = f"Processing complete.{suffix}"
            self.status.alert_type = "success"
            self._render_all()
        except Exception as exc:
            self.status.object = f"Error during processing: {exc}"
            self.status.alert_type = "danger"

    def _render_all(self):
        self._render_upload()
        self._render_process()
        self._render_analysis()
        self._render_download()

    def _data_table(self, df: pd.DataFrame, height: int = 320):
        return pn.widgets.Tabulator(df, pagination="remote", page_size=10, height=height, disabled=True)

    def _metrics(self, items):
        return pn.Row(*[pn.pane.HTML(f"<div class='metric'><b>{value}</b><span>{label}</span></div>") for label, value in items])

    def _render_upload(self):
        items = [
            pn.pane.Markdown("## Upload Data"),
            pn.Row(self.file_input, self.demo_button),
            self.status,
        ]
        if self.df_raw is not None:
            df = self.df_raw
            missing = df.isnull().sum()
            missing_df = pd.DataFrame(
                {
                    "Column": missing.index,
                    "Missing Count": missing.values,
                    "Missing %": (missing / len(df) * 100).round(2).values,
                }
            ).sort_values("Missing Count", ascending=False)
            col_info = pd.DataFrame(
                {
                    "Column": df.columns,
                    "Type": df.dtypes.astype(str).values,
                    "Non-Null Count": df.notna().sum().values,
                    "Unique Values": df.nunique().values,
                }
            )
            numeric_cols = df.select_dtypes(include=["number"]).columns
            stats_table = df[numeric_cols].describe() if len(numeric_cols) else pd.DataFrame({"Message": ["No numeric columns found"]})
            tabs = pn.Tabs(
                ("Preview", pn.Column(self.rows_slider, self._data_table(df.head(self.rows_slider.value)))),
                ("Statistics", self._data_table(stats_table)),
                ("Missing Values", self._data_table(missing_df[missing_df["Missing Count"] > 0] if missing_df["Missing Count"].sum() else pd.DataFrame({"Message": ["No missing values found"]}))),
                ("Columns", self._data_table(col_info)),
            )
            items.extend([self._metrics([("Rows", f"{len(df):,}"), ("Columns", len(df.columns)), ("File", self.uploaded_filename)]), tabs])
        self.upload_view.objects = items

    def _render_process(self):
        items = [
            pn.pane.Markdown("## Process Data"),
            pn.pane.Markdown("Clean the dataset and classify subjects by sex using numeric identifiers."),
        ]
        if self.df_raw is None:
            items.append(pn.pane.Alert("No data loaded yet.", alert_type="warning"))
        else:
            items.extend([
                pn.Row(self.sex_col, self.mode),
                pn.Row(self.threshold, self.male_ids, self.female_ids),
                self.process_button,
            ])
        if self.df_processed is not None:
            counts = self.df_processed["Sex"].value_counts()
            items.extend([
                self._metrics([
                    ("Total Rows", f"{len(self.df_processed):,}"),
                    ("Males", counts.get("Male", 0)),
                    ("Females", counts.get("Female", 0)),
                ]),
                self._data_table(self.df_processed.head(20)),
            ])
        self.process_view.objects = items

    def _sex_distribution_figure(self):
        counts = self.df_processed["Sex"].value_counts()
        fig, axes = plt.subplots(1, 2, figsize=(10, 4))
        colors = ["#3498db", "#e74c3c", "#9ca7bf"]
        axes[0].pie(counts.values, labels=counts.index, autopct="%1.1f%%", colors=colors[: len(counts)], startangle=90)
        axes[0].set_title("Sex Distribution")
        counts.plot(kind="bar", ax=axes[1], color=colors[: len(counts)], edgecolor="black")
        axes[1].set_title("Sex Counts")
        axes[1].set_xlabel("Sex")
        axes[1].set_ylabel("Count")
        axes[1].tick_params(axis="x", rotation=0)
        fig.tight_layout()
        return fig

    def _comparison_figure(self, selected_var: str):
        fig, axes = plt.subplots(1, 3, figsize=(15, 4))
        self.df_processed.boxplot(column=selected_var, by="Sex", ax=axes[0])
        axes[0].set_title(f"{selected_var} by Sex")
        axes[0].set_xlabel("Sex")
        axes[0].set_ylabel(selected_var)
        male = self.df_processed[self.df_processed["Sex"] == "Male"][selected_var].dropna()
        female = self.df_processed[self.df_processed["Sex"] == "Female"][selected_var].dropna()
        axes[1].violinplot([male, female], positions=[1, 2], showmeans=True, showmedians=True)
        axes[1].set_xticks([1, 2])
        axes[1].set_xticklabels(["Male", "Female"])
        axes[1].set_title("Distribution by Sex")
        axes[2].hist(male, bins=20, alpha=0.6, label="Male", color="#3498db", edgecolor="black")
        axes[2].hist(female, bins=20, alpha=0.6, label="Female", color="#e74c3c", edgecolor="black")
        axes[2].set_title("Distribution Overlay")
        axes[2].legend()
        fig.suptitle("")
        fig.tight_layout()
        return fig

    def _comparison_table(self, selected_vars: list[str]) -> pd.DataFrame:
        rows = []
        for var in selected_vars:
            male_vals = self.df_processed[self.df_processed["Sex"] == "Male"][var].dropna()
            female_vals = self.df_processed[self.df_processed["Sex"] == "Female"][var].dropna()
            male_mean = male_vals.mean()
            female_mean = female_vals.mean()
            t_stat, p_value, n_male, n_female = ttest_for_groups(self.df_processed, var)
            rows.append(
                {
                    "Variable": var,
                    "Male Mean": male_mean,
                    "Female Mean": female_mean,
                    "Difference": male_mean - female_mean,
                    "Diff %": safe_percent_diff(male_mean, female_mean),
                    "Male n": n_male,
                    "Female n": n_female,
                    "t-stat": t_stat,
                    "p-value": p_value,
                    "Cohen's d": effect_size_cohens_d(male_vals, female_vals),
                    "Significance": "Insufficient data" if pd.isna(p_value) else ("Significant" if p_value < 0.05 else "Not Significant"),
                }
            )
        return pd.DataFrame(rows).round(4)

    def _render_analysis(self):
        items = [pn.pane.Markdown("## Sex Analysis")]
        if self.df_processed is None or "Sex" not in self.df_processed.columns:
            items.append(pn.pane.Alert("Run the processing pipeline before analysis.", alert_type="warning"))
            self.analysis_view.objects = items
            return
        counts = self.df_processed["Sex"].value_counts()
        items.extend([
            self._metrics([
                ("Total Samples", f"{len(self.df_processed):,}"),
                ("Males", f"{counts.get('Male', 0)}"),
                ("Females", f"{counts.get('Female', 0)}"),
            ]),
            pn.pane.Matplotlib(self._sex_distribution_figure(), tight=True),
        ])
        numeric_cols = self.df_processed.select_dtypes(include=["number"]).columns.tolist()
        if numeric_cols:
            if self.compare_var.value not in numeric_cols:
                self.compare_var.value = numeric_cols[0]
            selected = self.compare_var.value
            t_stat, p_value, n_male, n_female = ttest_for_groups(self.df_processed, selected)
            stats_df = self.df_processed.groupby("Sex")[selected].agg(["count", "mean", "std", "min", "median", "max"]).round(4)
            test_label = "Insufficient data"
            if not pd.isna(t_stat):
                test_label = f"Welch t={t_stat:.4f}, p={p_value:.4f} (Male n={n_male}, Female n={n_female})"
            items.extend([
                pn.Row(self.compare_var),
                pn.pane.Matplotlib(self._comparison_figure(selected), tight=True),
                pn.pane.Alert(test_label, alert_type="success" if not pd.isna(p_value) and p_value < 0.05 else "info"),
                pn.pane.Markdown("### Summary Statistics by Sex"),
                self._data_table(stats_df),
                pn.Row(self.multi_vars),
                pn.pane.Markdown("### Multi-variable Summary"),
                self._data_table(self._comparison_table(list(self.multi_vars.value or []))),
            ])
        self.analysis_view.objects = items

    def _csv_download(self):
        if self.df_processed is None:
            return io.StringIO("")
        buffer = io.StringIO()
        self.df_processed.to_csv(buffer, index=False)
        buffer.seek(0)
        return buffer

    def _excel_download(self):
        if self.df_processed is None:
            return io.BytesIO()
        return io.BytesIO(
            build_excel_export(
                self.df_processed,
                animal_col=self.sex_col_used,
                threshold=self.sex_threshold_used,
                rebuild_sex_labels=self.sex_rebuild_from_threshold,
            )
        )

    def _render_download(self):
        items = [pn.pane.Markdown("## Download Results")]
        if self.df_processed is None:
            items.append(pn.pane.Alert("No processed data available yet.", alert_type="warning"))
        else:
            stem = Path(self.uploaded_filename).stem
            methods_text = (
                "Data were imported and processed using the NeuroPipe-SABV Panel application "
                "(Anderson & Devan, 2026). Subjects were classified by sex based on numerical "
                f"identifiers extracted from '{self.sex_col_used}'. Missing values were handled "
                "during group comparisons. Statistical differences between male and female cohorts "
                "were assessed using Welch's independent samples t-tests."
            )
            csv_size = len(self.df_processed.to_csv(index=False).encode("utf-8")) / 1024
            items.extend([
                pn.Row(
                    pn.widgets.FileDownload(
                        label="Download CSV",
                        filename=f"{stem}_processed.csv",
                        callback=self._csv_download,
                        button_type="primary",
                    ),
                    pn.widgets.FileDownload(
                        label="Download Excel",
                        filename=f"{stem}_processed.xlsx",
                        callback=self._excel_download,
                        button_type="success",
                    ),
                ),
                self._metrics([
                    ("Rows", f"{len(self.df_processed):,}"),
                    ("Columns", len(self.df_processed.columns)),
                    ("CSV Size", f"{csv_size:.1f} KB"),
                ]),
                pn.pane.Markdown("### Methods Section Draft"),
                pn.pane.Str(methods_text, styles={"white-space": "pre-wrap"}),
            ])
        self.download_view.objects = items

    def view(self):
        hero = pn.pane.HTML(
            """
            <section class="hero">
              <p class="muted"><strong>NIH</strong> = National Institutes of Health | <strong>SABV</strong> = Sex as a Biological Variable</p>
              <h1>NeuroPipe-SABV</h1>
              <p>Open-source Panel workflow for preclinical behavioral data with attention to sex as a biological variable.</p>
            </section>
            """
        )
        tabs = pn.Tabs(
            ("Upload", self.upload_view),
            ("Process", self.process_view),
            ("Analyze", self.analysis_view),
            ("Download", self.download_view),
            dynamic=True,
        )
        return pn.Column(hero, tabs, max_width=1180, margin=(20, 32))


app = NeuroPipePanel()
app.view().servable(title="NeuroPipe-SABV")

